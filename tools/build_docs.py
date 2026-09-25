"""
Build the lab record Word documents (CD05-104.docx ... CD08-104.docx) from
LabRecordTemplate.docx.  The code is read from the assignment folders and the
screenshots (output.png, output2.png ...) are taken from the same folders, so if you
replace a screenshot with your own, just run this script again.

Usage :  python tools/build_docs.py            (all records)
         python tools/build_docs.py 6 8        (only Assignment6 and Assignment8)
"""
import copy
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import NAME, RECORDS, REG_NO  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "tools", "LabRecordTemplate.docx")

BODY_PT = 11
CODE_PT = 8.5
CODE_FONT = "Consolas"
MAX_IMG_W_CM = 16.2
MAX_IMG_H_CM = 20.5


# ------------------------------------------------------------------ xml helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


def set_run_font(run, name=None, size=None, bold=None, italic=None, noproof=False):
    """Elements are added in the order Word expects: rFonts, b, i, noProof, sz."""
    rpr = run._r.get_or_add_rPr()
    if name:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = el("w:rFonts")
            rpr.insert(0, rfonts)
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            rfonts.set(qn("w:" + a), name)
    if noproof:
        rpr.append(el("w:noProof"))
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if size:
        run.font.size = Pt(size)


def spacing(p, before=0, after=6, line=None, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True


# ---------------------------------------------------------------- numbering (bullets)
def add_bullet_numbering(doc):
    """The template has no bullet list, so add one (returns its numId)."""
    numbering = doc.part.numbering_part.element
    abstract_id, num_id = 90, 90
    abstract = el("w:abstractNum", w_abstractNumId=abstract_id)
    abstract.append(el("w:multiLevelType", w_val="hybridMultilevel"))
    lvl = el("w:lvl", w_ilvl=0)
    lvl.append(el("w:start", w_val=1))
    lvl.append(el("w:numFmt", w_val="bullet"))
    lvl.append(el("w:lvlText", w_val="•"))
    lvl.append(el("w:lvlJc", w_val="left"))
    ppr = el("w:pPr")
    ppr.append(el("w:ind", w_left=720, w_hanging=360))
    lvl.append(ppr)
    rpr = el("w:rPr")
    rpr.append(el("w:rFonts", w_ascii="Calibri", w_hAnsi="Calibri", w_cs="Calibri"))
    lvl.append(rpr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = el("w:num", w_numId=num_id)
    num.append(el("w:abstractNumId", w_val=abstract_id))
    numbering.append(num)
    return num_id


# ------------------------------------------------------------------- block builders
def heading(doc, text, before=10):
    p = doc.add_paragraph()
    spacing(p, before=before, after=4, keep_next=True)
    set_run_font(p.add_run(text), size=12, bold=True)
    return p


def paragraph(doc, text, italic=False, size=BODY_PT, after=6, align=None):
    p = doc.add_paragraph()
    spacing(p, after=after)
    if align:
        p.alignment = align
    set_run_font(p.add_run(text), size=size, italic=italic)
    return p


def bullet(doc, text, num_id):
    p = doc.add_paragraph(style="List Paragraph")
    ppr = p._p.get_or_add_pPr()
    numpr = el("w:numPr")
    numpr.append(el("w:ilvl", w_val=0))
    numpr.append(el("w:numId", w_val=num_id))
    ppr.append(numpr)
    spacing(p, after=3)
    p.paragraph_format.line_spacing = 1.1
    set_run_font(p.add_run(text), size=BODY_PT)
    return p


def code_block(doc, text):
    """Monospace, shaded box with a thin border (adjacent paragraphs merge into one box)."""
    lines = text.replace("\t", "    ").rstrip("\n").split("\n")
    for k, ln in enumerate(lines):
        p = doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        if len(lines) <= 25 and k < len(lines) - 1:   # short blocks stay on one page
            ppr.append(el("w:keepNext"))
        bdr = el("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            bdr.append(el("w:" + side, w_val="single", w_sz=4, w_space=2, w_color="BFBFBF"))
        ppr.append(bdr)
        ppr.append(el("w:shd", w_val="clear", w_color="auto", w_fill="F5F5F5"))
        ppr.append(el("w:spacing", w_before=0, w_after=0, w_line=240, w_lineRule="auto"))
        ppr.append(el("w:ind", w_left=113, w_right=113))
        mark = el("w:rPr")                        # size of the paragraph mark (empty lines)
        mark.append(el("w:rFonts", w_ascii=CODE_FONT, w_hAnsi=CODE_FONT, w_cs=CODE_FONT))
        mark.append(el("w:noProof"))
        mark.append(el("w:sz", w_val=int(CODE_PT * 2)))
        ppr.append(mark)
        if ln:
            set_run_font(p.add_run(ln), name=CODE_FONT, size=CODE_PT, noproof=True)
    gap = doc.add_paragraph()                     # small gap after the box
    spacing(gap, after=0)
    gap.paragraph_format.line_spacing = 0.6


def file_label(doc, text):
    p = doc.add_paragraph()
    spacing(p, before=6, after=3, keep_next=True)
    set_run_font(p.add_run(text), size=BODY_PT, bold=True)


def picture(doc, path, caption):
    with Image.open(path) as im:
        w_px, h_px = im.size
    width = MAX_IMG_W_CM
    height = width * h_px / w_px
    if height > MAX_IMG_H_CM:
        height = MAX_IMG_H_CM
        width = height * w_px / h_px
    doc.add_picture(path, width=Cm(width), height=Cm(height))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, before=4, after=2, keep_next=True)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(c, after=8)
    set_run_font(c.add_run(caption), size=9.5, italic=True)


# ------------------------------------------------------------------- header table
def fill_cell(cell, label, value):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(5)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    ppr = p._p.pPr
    if ppr is not None:
        jc = ppr.find(qn("w:jc"))
        if jc is not None:
            ppr.remove(jc)
    r1 = p.add_run(label)
    r1.bold = True
    p.add_run(" " + value)


def fill_header_table(doc, number, date):
    tbl = doc.tables[0]
    first = tbl.rows[0]._tr
    clone = copy.deepcopy(first)
    for node in clone.iter():                     # no duplicate paragraph ids
        for attr in list(node.attrib):
            if attr.endswith("}paraId") or attr.endswith("}textId"):
                del node.attrib[attr]
    first.addprevious(clone)                      # new first row: Name / Reg. No
    rows = tbl.rows
    fill_cell(rows[0].cells[0], "Name:", NAME)
    fill_cell(rows[0].cells[1], "Reg. No:", REG_NO)
    fill_cell(rows[1].cells[0], "Exp. No:", str(number))
    fill_cell(rows[1].cells[1], "Date:", date)
    for row in rows:                              # bold labels only, like the earlier records
        for cell in row.cells:
            for r in cell.paragraphs[0].runs[1:]:
                r.bold = False


# --------------------------------------------------------------------- one record
def build(number):
    rec = RECORDS[number]
    folder = os.path.join(ROOT, "Assignment%d" % number)
    out = os.path.join(folder, "CD%02d-104.docx" % number)

    doc = Document(TEMPLATE)
    body = doc.element.body
    for child in list(body):                      # keep the table and the section settings
        if child.tag == qn("w:p"):
            body.remove(child)
    bullets = add_bullet_numbering(doc)
    fill_header_table(doc, number, rec["date"])

    doc.add_paragraph()                           # blank line before the title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(t, before=6, after=8)
    set_run_font(t.add_run(rec["title"]), size=14, bold=True)

    heading(doc, "Aim:", before=4)
    paragraph(doc, rec["aim"])

    heading(doc, "Tools Used:")
    for tool in rec["tools"]:
        bullet(doc, tool, bullets)

    heading(doc, "Code:")
    for label, fname in rec["code"]:
        file_label(doc, label)
        with open(os.path.join(folder, fname), encoding="utf-8") as f:
            code_block(doc, f.read())

    for label, fname, note in rec["inputs"]:
        file_label(doc, label)
        with open(os.path.join(folder, fname), encoding="utf-8") as f:
            code_block(doc, f.read())
        if note:
            paragraph(doc, note, italic=True, size=10)

    file_label(doc, "Commands Used")
    code_block(doc, "\n".join(rec["commands"]))

    heading(doc, "Output:")
    for image, caption in rec["outputs"]:
        picture(doc, os.path.join(folder, image), caption)
    for text in rec["notes"]:
        paragraph(doc, text, size=10.5)

    if rec.get("extra"):
        heading(doc, "Additional Test Cases:")
        for image, caption in rec["extra"]:
            picture(doc, os.path.join(folder, image), caption)
        for text in rec.get("extra_notes", []):
            paragraph(doc, text, size=10.5)

    heading(doc, "Learning Outcomes:", before=12)
    for item in rec["learning"]:
        bullet(doc, item, bullets)

    cp = doc.core_properties
    cp.author = NAME
    cp.title = "Compiler Design Lab - Experiment %d" % number
    cp.last_modified_by = NAME
    doc.save(out)
    print("written", os.path.relpath(out, ROOT))


if __name__ == "__main__":
    wanted = [int(a) for a in sys.argv[1:]] or sorted(RECORDS)
    for n in wanted:
        build(n)
